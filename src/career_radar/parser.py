"""
HTML announcement content and multi-format attachment parser (XLSX, DOCX, PDF).
Implements Issue #10 deterministic extraction without business rules.
"""

import os
from pathlib import Path
import re
from typing import Any, Dict, List, Optional
from urllib.parse import urljoin
from bs4 import BeautifulSoup
import docx
import openpyxl
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".xlsx", ".docx", ".pdf"}
LEGACY_EXTENSIONS = {".xls", ".doc"}
ALL_DISCOVERY_EXTENSIONS = SUPPORTED_EXTENSIONS | LEGACY_EXTENSIONS


class HTMLAnnouncementParser:
    """
    Extracts announcement metadata, text, and discovered attachment links from HTML.
    """

    def parse(self, html_content: str, base_url: str = "") -> Dict[str, Any]:
        soup = BeautifulSoup(html_content, "html.parser")

        # 1. Title discovery
        title = ""
        title_el = (
            soup.find("h1")
            or soup.find("div", class_=re.compile(r"title|article-title|header", re.I))
            or soup.find("title")
        )
        if title_el:
            title = title_el.get_text(strip=True)

        # 2. Body text extraction
        body_text = soup.get_text(separator="\n", strip=True)

        # 3. Headings extraction
        headings = []
        for h_tag in soup.find_all(re.compile(r"^h[1-6]$", re.I)):
            h_text = h_tag.get_text(strip=True)
            if h_text:
                headings.append({"level": h_tag.name.lower(), "text": h_text})

        # 4. Links extraction
        links = []
        for a_tag in soup.find_all("a", href=True):
            href = a_tag["href"].strip()
            l_text = a_tag.get_text(strip=True)
            full_url = urljoin(base_url, href) if base_url else href
            links.append({"text": l_text, "url": full_url})

        # 5. HTML Table extraction
        tables = []
        for t_idx, table_el in enumerate(soup.find_all("table")):
            rows_data = []
            for tr in table_el.find_all("tr"):
                cells = [td.get_text(separator=" ", strip=True) for td in tr.find_all(["th", "td"])]
                if any(cells):
                    rows_data.append(cells)
            if not rows_data:
                continue

            header_idx = 0
            headers = []
            for idx, r in enumerate(rows_data):
                non_empty = [c for c in r if c]
                if len(non_empty) >= 2:
                    header_idx = idx
                    headers = [c if c else f"col_{i}" for i, c in enumerate(r)]
                    break
            if not headers:
                continue

            parsed_rows = []
            for row_idx, r in enumerate(rows_data[header_idx + 1:], start=header_idx + 2):
                if not any(r):
                    continue
                cell_dict = {}
                for col_idx, cell_val in enumerate(r):
                    if col_idx < len(headers):
                        cell_dict[headers[col_idx]] = cell_val
                parsed_rows.append({"row_index": row_idx, "cells": cell_dict})

            tables.append({
                "file_type": "html_table",
                "table_index": t_idx,
                "status": "success",
                "headers": headers,
                "rows": parsed_rows,
            })

        # 6. Discovered Attachment links
        attachments = []
        for a_tag in soup.find_all("a", href=True):
            href = a_tag["href"].strip()
            link_text = a_tag.get_text(strip=True) or os.path.basename(href)
            link_title = a_tag.get("title", "").strip()
            full_url = urljoin(base_url, href) if base_url else href

            # Check extension in href path, link text, or link title
            detected_ext = ""
            for ext in ALL_DISCOVERY_EXTENSIONS:
                if (
                    href.lower().endswith(ext)
                    or href.lower().split("?")[0].endswith(ext)
                    or link_text.lower().endswith(ext)
                    or link_title.lower().endswith(ext)
                ):
                    detected_ext = ext
                    break

            # Also check if it's an attachment download endpoint with extension in text/title
            if not detected_ext and any(k in href.lower() for k in ["download", "attach", "wbfileid"]):
                for ext in ALL_DISCOVERY_EXTENSIONS:
                    if ext in link_text.lower() or ext in link_title.lower():
                        detected_ext = ext
                        break

            if detected_ext:
                is_supported = detected_ext in SUPPORTED_EXTENSIONS
                attachments.append({
                    "name": link_text,
                    "url": full_url,
                    "extension": detected_ext,
                    "supported": is_supported,
                    "status": "supported" if is_supported else "unsupported_legacy_format",
                })

        return {
            "title": title,
            "body_text": body_text,
            "headings": headings,
            "links": links,
            "tables": tables,
            "attachments": attachments,
        }


class AttachmentParser:
    """
    Extracts tabular rows and text from XLSX, DOCX, and text-native PDF attachments.
    Explicitly refuses unsupported legacy formats (.xls, .doc) with Evidence Gap markers.
    Maintains row indices, column headers, and cell content with zero business filtering.
    """

    def parse_file(self, file_path: Path) -> List[Dict[str, Any]]:
        ext = file_path.suffix.lower()
        if ext == ".xlsx":
            return self._parse_xlsx(file_path)
        elif ext == ".docx":
            return self._parse_docx(file_path)
        elif ext == ".pdf":
            return self._parse_pdf(file_path)
        elif ext in LEGACY_EXTENSIONS:
            return [
                {
                    "file_type": ext.lstrip("."),
                    "file_name": file_path.name,
                    "status": "unsupported_legacy_format",
                    "error": (
                        f"Legacy format '{ext}' is not supported by mechanical parser. "
                        "Issue #10 officially supports .xlsx, .docx, and text-native .pdf."
                    ),
                    "rows": [],
                }
            ]
        return [
            {
                "file_type": ext.lstrip(".") or "unknown",
                "file_name": file_path.name,
                "status": "unsupported_format",
                "rows": [],
            }
        ]

    def _parse_xlsx(self, file_path: Path) -> List[Dict[str, Any]]:
        tables = []
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                rows_data = list(sheet.iter_rows(values_only=True))
                if not rows_data:
                    continue

                # Find first non-empty row with at least 2 headers
                header_idx = 0
                headers = []
                for idx, r in enumerate(rows_data):
                    non_empty = [str(c).strip() for c in r if c is not None and str(c).strip()]
                    if len(non_empty) >= 2:
                        header_idx = idx
                        headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(r)]
                        break

                if not headers:
                    continue

                parsed_rows = []
                for row_idx, r in enumerate(rows_data[header_idx + 1:], start=header_idx + 2):
                    if not any(c is not None and str(c).strip() for c in r):
                        continue
                    cell_dict = {}
                    for col_idx, cell_val in enumerate(r):
                        if col_idx < len(headers):
                            h_name = headers[col_idx]
                            cell_dict[h_name] = str(cell_val).strip() if cell_val is not None else ""
                    parsed_rows.append({"row_index": row_idx, "cells": cell_dict})

                tables.append({
                    "file_type": "xlsx",
                    "file_name": file_path.name,
                    "sheet_name": sheet.title,
                    "status": "success",
                    "headers": headers,
                    "rows": parsed_rows,
                })
        except Exception as e:
            tables.append({
                "file_type": "xlsx",
                "file_name": file_path.name,
                "status": "error",
                "error": str(e),
                "rows": [],
            })
        return tables

    def _parse_docx(self, file_path: Path) -> List[Dict[str, Any]]:
        tables = []
        try:
            doc = docx.Document(file_path)
            for t_idx, table in enumerate(doc.tables):
                rows = table.rows
                if not rows:
                    continue
                headers = [cell.text.strip() for cell in rows[0].cells]
                parsed_rows = []
                for r_idx, row in enumerate(rows[1:], start=2):
                    cell_dict = {}
                    for c_idx, cell in enumerate(row.cells):
                        if c_idx < len(headers):
                            cell_dict[headers[c_idx]] = cell.text.strip()
                    parsed_rows.append({"row_index": r_idx, "cells": cell_dict})

                tables.append({
                    "file_type": "docx",
                    "file_name": file_path.name,
                    "table_index": t_idx,
                    "status": "success",
                    "headers": headers,
                    "rows": parsed_rows,
                })
        except Exception as e:
            tables.append({
                "file_type": "docx",
                "file_name": file_path.name,
                "status": "error",
                "error": str(e),
                "rows": [],
            })
        return tables

    def _parse_pdf(self, file_path: Path) -> List[Dict[str, Any]]:
        tables = []
        try:
            reader = PdfReader(file_path)
            pages_text = []
            for p_idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                pages_text.append({"page": p_idx + 1, "text": text.strip()})
            tables.append({
                "file_type": "pdf",
                "file_name": file_path.name,
                "status": "success",
                "pages": pages_text,
                "rows": [],
            })
        except Exception as e:
            tables.append({
                "file_type": "pdf",
                "file_name": file_path.name,
                "status": "error",
                "error": str(e),
                "rows": [],
            })
        return tables
